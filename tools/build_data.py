#!/usr/bin/env python3
"""
Build data/products.json (and optimized product images) from the Excel pricelist.

The Excel file is the single source of truth. Re-run this script whenever the
pricelist or the product photos change, then commit the results.

Usage (from the project root):
    pip install openpyxl pillow
    python tools/build_data.py --excel "source/Current Pricelist - HPNB.xlsx" --images "path/to/Image assets"

Image matching rule
-------------------
Images in the --images folder must be named "<exact Model from Excel> - <Angle>.<ext>",
e.g. "HP OmniBook 5 AI PC 14-kf0002TU - Front.png". A file is only assigned to a
product when its name starts with that product's exact Model text followed by " - ".
Files that do not match any model are reported and skipped (never guessed).
"""
import argparse
import hashlib
import json
import os
import re
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import changes  # noqa: E402  (tracks new models / price changes for the website badges)

try:
    import openpyxl
except ImportError:
    sys.exit("Please install openpyxl:  pip install openpyxl pillow")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_OUT = os.path.join(ROOT, "assets", "images", "products")
DATA_OUT = os.path.join(ROOT, "data")

# Preferred display order for known photo angles (unknown angles go last, alphabetically)
ANGLE_ORDER = ["front", "front right", "front left", "rear left", "rear right", "rear",
               "left profile", "right profile", "side profiles", "stacked profile", "open", "lifestyle"]
# HP spec-sheet style angle names -> the names used on the site
ANGLE_ALIAS = {"center": "Front", "centre": "Front", "front left": "Front Left", "rear left": "Rear Left",
               "sides": "Side Profiles"}

# Friendly labels + short descriptions for the Type codes used in the Excel.
# New Type codes found in the Excel are added automatically with a neutral description.
TYPE_INFO_DEFAULTS = {
    "HP 15": ("HP Laptop 15", "Everyday 15.6-inch laptops for home, school and small-office productivity."),
    "OB3": ("OmniBook 3", "Designed for everyday productivity with a balance of performance, portability, and value."),
    "OB5": ("OmniBook 5", "Thin-and-light 14-inch laptops for mainstream productivity and on-the-go work."),
    "OB5 FLIP": ("OmniBook 5 Flip", "Convertible 2-in-1 laptops that switch between laptop and tablet modes for work, study and creative tasks."),
    "OBX": ("OmniBook X", "Premium OmniBook X Flip convertible AI PCs for professionals who want performance in a versatile 2-in-1 design."),
    "Omnibook 7": ("OmniBook 7", "Premium 14-inch OmniBook 7 AI PCs with OLED displays for professionals and creators."),
    "Probook": ("ProBook", "Business laptops for mainstream business productivity, with HP Wolf Security built in."),
}

# order of the model tabs (entry level first, premium last); unknown types go to the end in Excel order
TYPE_RANK = ["HP 15", "Omnibook 3", "Omnibook 5", "Omnibook 5 Flip", "Omnibook 7", "Omnibook X", "Omnibook X Flip", "Probook"]


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = re.sub(r"\s+", " ", v).strip()
        return v or None
    return v


def slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def parse_specs(spec):
    """Split the pipe-delimited Specs string into labelled fields using keyword detection.
    Nothing is invented: a field is only set if a matching segment exists."""
    out = {}
    if not spec:
        return out
    parts = [p.strip() for p in spec.split("|") if p.strip()]
    rest = []
    for i, p in enumerate(parts):
        pl = p.lower()
        if i == 0:
            out["series"] = p
            m = re.search(r"(clamshell|x360|detachable)", pl)
            if m:
                out["formFactor"] = {"clamshell": "Clamshell", "x360": "x360 Convertible (2-in-1)",
                                     "detachable": "Detachable"}[m.group(1)]
            continue
        if "cpu" not in out and re.search(r"\b(ryzen|core|ultra|celeron|pentium|snapdragon|athlon|intel|amd)\b", pl) \
                and "graphics" not in pl:
            out["cpu"] = p
        elif "ram" not in out and re.search(r"\d+\s?gb\b.*(ddr|lpddr)", pl):
            out["ram"] = p
        elif "storage" not in out and re.search(r"\d+\s?(gb|tb)\b.*(pcie|ssd|emmc|nvme|hdd)", pl):
            out["storage"] = p
        elif "graphics" not in out and "graphics" in pl:
            out["graphics"] = p
        elif "display" not in out and re.search(r"\d{2}\.\d", pl) and re.search(r"(fhd|2k|3k|4k|hd|nits|oled|ips)", pl):
            out["display"] = p
        elif "os" not in out and re.search(r"\b(ost|w11|win|windows|dos|chrome)", pl):
            out["os"] = p
        elif "colorCamera" not in out and " - " in p:
            color, cam = p.split(" - ", 1)
            out["color"] = color.strip()
            out["camera"] = cam.strip()
            out["colorCamera"] = p
        elif "adapter" not in out and re.search(r"\bwatt\b|\bw\b adapter", pl):
            out["adapter"] = p
        elif "warranty" not in out and (re.fullmatch(r"\d\s*-\s*\d\s*-\s*\d", p) or pl.startswith("warr")):
            out["warranty"] = p
        elif "carePack" not in out and re.search(r"(care|svc|bundle|onsite)", pl):
            out["carePack"] = p
        else:
            rest.append(p)
    out.pop("colorCamera", None)
    if rest:
        out["other"] = rest
    return out


def cpu_brand(cpu):
    if not cpu:
        return None
    c = cpu.lower()
    if re.search(r"\b(ryzen|athlon|amd)\b", c):
        return "AMD"
    if re.search(r"\b(core|ultra|celeron|pentium|intel|xeon)\b", c):
        return "Intel"
    if "snapdragon" in c:
        return "Qualcomm"
    return None


def display_short(d):
    if not d:
        return None
    touch = d.lower().startswith("touch")
    d2 = re.sub(r"^touch/", "", d, flags=re.I)
    size = re.search(r"(\d{2}\.\d)", d2)
    res = re.search(r"\b(FHD|2K|3K|4K|2\.8K|HD\+?)\b(?:\s*\((\d+x\d+)\))?", d2)
    bits = []
    if size:
        bits.append(size.group(1) + '"')
    if res:
        bits.append(res.group(1) + (f" ({res.group(2)})" if res.group(2) else ""))
    if re.search(r"\boled\b", d2, re.I):
        bits.append("OLED")
    if touch:
        bits.append("Touch")
    return " ".join(bits) if bits else d


def sku_from_model(model):
    m = re.search(r"(\d{2}-[A-Za-z]{2}\d{4}[A-Za-z]{2})\s*$", model) or \
        re.search(r"\b(\d{2}-[A-Za-z]{1,2}\d{4}[A-Za-z]{2})\b", model)
    return m.group(1) if m else model.split()[-1]


def product_sku(model, part_no):
    """SKU shown on the site: the HP SKU in the model name (e.g. 14-kb0100TU); for models named
    without one (e.g. business 'HP Probook 4 G1I 14 AI ...') the HP part number is used."""
    m = re.search(r"(\d{2}-[A-Za-z]{2}\d{4}[A-Za-z]{2})\s*$", model) or \
        re.search(r"\b(\d{2}-[A-Za-z]{1,2}\d{4}[A-Za-z]{2})\b", model)
    if m:
        return m.group(1)
    return part_no if part_no else model.split()[-1]


PART_RE = re.compile(r"^[A-Z0-9]{5,8}(#[A-Z0-9]{3})?$")


# ---------------------------------------------------------------------------
# HP technical spec sheets (source/specsheets/<PartNo>_*.xlsx, exported from HP)
# ---------------------------------------------------------------------------
SURE_INFO = {
    # short sales explanations; requirements come from the spec sheet's own footnotes
    "HP Sure Start": "Self-healing BIOS: checks the BIOS every time the PC starts and automatically restores a clean copy if it has been attacked or corrupted.",
    "HP Sure Run": "Keeps critical security processes running: if malware tries to stop them, Sure Run restarts them automatically.",
    "HP Sure Sense": "AI-based malware protection that uses deep learning to stop new, never-seen-before malware and ransomware.",
    "HP Sure Click": "Opens risky web pages and files in isolated micro virtual machines, so any malware is contained and removed when the tab or file is closed.",
    "HP Sure Recover": "Reinstalls the Windows image over a network connection, even if the drive was wiped - no USB stick or IT visit needed.",
    "HP Sure Admin": "Lets IT manage BIOS settings with secure digital certificates instead of a shared BIOS password; local access is approved with a smartphone app.",
}


def read_specsheet(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    kv = {}
    for r in ws.iter_rows(values_only=True):
        if not r or r[0] is None or len(r) < 2 or r[1] is None:
            continue
        k, v = str(r[0]).strip(), clean(str(r[1]))
        if k in kv:   # repeated keys (e.g. 'Battery Type'): keep the most descriptive value
            if len(v) > len(kv[k]):
                kv[k] = v
        else:
            kv[k] = v
    return kv


def specsheet_details(kv, fname):
    g = kv.get
    notes = {}
    for k, v in kv.items():
        m = re.match(r"Tech Spec Footnote \d+", k)
        n = re.match(r"\[(\d+)\]\s*(.*)", v or "")
        if m and n:
            notes[n.group(1)] = n.group(2).strip()

    def fn(key):
        ref = g(key) or ""
        return [notes[x] for x in re.findall(r"\d+", ref) if x in notes]

    mem = g("Standard Memory Description 01")
    slots = g("Memory Slots")
    layout = g("Memory Layout (Slots & Size)")
    battery = g("Battery Type")
    fast = "Fast charging 50% in 30 minutes" if "fast charging 50% in 30 minutes" in (g("Battery Life Note") or "").lower() else None
    mil = re.search(r"MIL-STD[- ]?\d+\w*", g("Certifications and Compliances") or "")
    kb = ", ".join(x for x in [g("Keyboard")] if x)
    rows = [
        ("Processor", g("Processor Name")),
        ("AI engine (NPU)", g("Neural Processing Unit (NPU)")),
        ("Memory", mem),
        ("Memory slots", f"{slots} ({layout} installed)" if slots and layout else slots),
        ("Storage", g("Hard Drive Description (01)")),
        ("Display", g("Display Description")),
        ("Graphics", g("Graphics (facet)") or g("GScardseg 01 Available Card 01")),
        ("Operating system", g("Operating System Software")),
        ("Camera", g("Integrated Camera")),
        ("Audio", g("Audio Features")),
        ("Wireless", g("Wireless Technologies")),
        ("Ethernet", g("Network Interface")),
        ("Ports", g("External I/O Ports")),
        ("Keyboard", kb or None),
        ("Fingerprint reader", "Yes" if g("Fingerprint Reader") else None),
        ("Battery", battery + (f" · {fast}" if fast else "") if battery else None),
        ("Power adapter", g("Power Supply Required")),
        ("Colour", g("Product Color Description (fixed)") or g("Color")),
        ("Weight", g("Weight (metric)")),
        ("Dimensions (W x D x H)", g("Product Dimensions W x D x H (metric)")),
        ("Durability", f"{mil.group(0)} tested" if mil else None),
        ("HP warranty (spec sheet)", g("Manufacturer Warranty")),
    ]
    sec = []
    for k in sorted((k for k in kv if re.fullmatch(r"Security Management\d+", k)), key=lambda k: int(k[19:])):
        sec.append({"name": kv[k], "notes": fn("Security Management Footnote Number" + k[19:])})
    sensors = [x.strip() for x in (g("Sensors") or "").split(";") if x.strip()]
    sure = []
    for s_ in sec:
        if s_["name"].startswith("HP Sure "):
            sure.append({"name": s_["name"], "what": SURE_INFO.get(s_["name"]), "notes": s_["notes"]})
    other_sec = [s_["name"] for s_ in sec if not s_["name"].startswith("HP Sure ")]
    for x in sensors:
        if x.startswith("HP Sure ") and x not in [s_["name"] for s_ in sure]:
            sure.append({"name": x + " (sensor)", "what": None, "notes": []})
    tm = lambda x: re.sub(r"[®™]", "", x or "").strip()
    short_cpu = None
    if g("Processor Family (Short)") and g("Processor Model"):
        short_cpu = f'{tm(g("Processor Family (Short)"))} {g("Processor Model")}' + (f' ({g("Processor Core")}C)' if g("Processor Core") else "") \
            + (f' {g("Processor NPU Tops")} NPU TOPs' if g("Processor NPU Tops") else "")
    short_ram = f'{g("Memory Size")}GB {g("Memory Type")} {g("Memory Speed")}' + (f' ({layout.replace(" GB", "GB")})' if layout else "") \
        if g("Memory Size") and g("Memory Type") and g("Memory Speed") else None
    short_sto = f'{g("Storage Capacity 01")}{g("Storage Capacity UOM 01")} {tm(g("Storage Interface 01"))} {tm(g("Storage Interface Protocol 01"))} {g("Storage Type 01")}' \
        if g("Storage Capacity 01") and g("Storage Type 01") else None
    disp_short = None
    if g("Display Size (Diagonal)"):
        disp_short = f'{g("Display Size (Diagonal)")}"' + (f' {g("Display Resolution Standard")}' if g("Display Resolution Standard") else "") \
            + (f' ({g("Display Resolution (Pixels)").replace(" ", "")})' if g("Display Resolution (Pixels)") else "")
    return {
        "name": g("Product name"),
        "productNumber": g("Product number"),
        "series": g("Series Name (facet)"),
        "usage": g("Usage (facet)"),
        "rows": [[k, v] for k, v in rows if v],
        "sure": sure,
        "security": other_sec,
        "sourceFile": fname,
        # fall-back values for the card / filters when the pricelist Specs text is short
        "_fill": {
            "cpu": short_cpu or g("Processor Name"), "ram": short_ram or mem, "storage": short_sto or g("Hard Drive Description (01)"),
            "graphics": g("Graphics (facet)"), "display": g("Display Description"), "displayShort": disp_short,
            "os": g("Operating System (facet)"), "color": g("Product Color Description (fixed)") or g("Color"),
            "camera": g("Integrated Camera"), "adapter": g("Power Supply Required"),
            "formFactor": g("Form Factor (facet)"), "series": g("Series Name (facet)"),
        },
    }


def load_specsheets(folder):
    out = {}
    if not folder or not os.path.isdir(folder):
        return out
    for f in sorted(os.listdir(folder)):
        if f.lower().endswith(".xlsx") and not f.startswith("~$"):
            kv = read_specsheet(os.path.join(folder, f))
            pn = (kv.get("Product number") or f.split("_")[0]).upper()
            out[pn] = specsheet_details(kv, f)
    return out


def read_excel(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header_row, headers = None, None
    for r in ws.iter_rows(min_row=1, max_row=20):
        vals = [clean(c.value) for c in r]
        low = [str(v).lower() if v is not None else "" for v in vals]
        if "model" in low and "srp" in low:
            header_row = r[0].row
            headers = vals
            break
    if not headers:
        sys.exit("Could not find a header row containing 'Model' and 'SRP'.")
    rows = []
    for r in ws.iter_rows(min_row=header_row + 1, values_only=True):
        rec = {}
        for h, v in zip(headers, r):
            if h is None:
                continue
            rec[str(h)] = clean(v)
        if rec.get("Model"):
            rows.append(rec)
    return headers, rows


def angle_key(a):
    a = a.lower()
    return (ANGLE_ORDER.index(a) if a in ANGLE_ORDER else len(ANGLE_ORDER), a)


def build_images(models, img_dir, max_px=1200, thumb_px=360, out_dir=None, url_prefix="assets/images/products", aliases=None):
    """Return {model: [ {angle, src, thumb}, ... ]} using exact model-name prefix matching.
    aliases: optional {part number: model} so "BQ5B1PT - Front.png" also matches its model."""
    from PIL import Image
    out_dir = out_dir or IMG_OUT
    os.makedirs(out_dir, exist_ok=True)
    result = {m: [] for m in models}
    unmatched = []
    files = sorted(f for f in os.listdir(img_dir) if f.lower().endswith((".png", ".jpg", ".jpeg", ".webp")))
    # longest model names first so "... 14-kb0100TU" can never swallow a longer name
    keys = {m: m for m in models}
    keys.update({k: m for k, m in (aliases or {}).items() if m in keys})
    ordered = sorted(keys, key=len, reverse=True)
    cache = {}  # identical photos shared by several models are only encoded once
    for f in files:
        stem = os.path.splitext(f)[0]
        key = next((m for m in ordered if stem.upper().startswith(m.upper() + " - ") if m != keys[m]), None) \
            or next((m for m in ordered if stem.startswith(m + " - ")), None)
        if not key:
            unmatched.append(f)
            continue
        owner = keys[key]
        angle = stem[len(key) + 3:].strip()
        angle = ANGLE_ALIAS.get(angle.lower(), angle)
        sku = key if key != owner else sku_from_model(owner)
        base = f"{slug(sku)}-{slug(angle)}"
        src_path = os.path.join(img_dir, f)
        out_main = os.path.join(out_dir, base + ".webp")
        out_thumb = os.path.join(out_dir, base + "-thumb.webp")
        with open(src_path, "rb") as fh:
            h = hashlib.md5(fh.read()).hexdigest()
        if h in cache:
            import shutil
            prev_main, prev_thumb, w, hh = cache[h]
            shutil.copyfile(prev_main, out_main)
            shutil.copyfile(prev_thumb, out_thumb)
            result[owner].append({"angle": angle, "src": f"{url_prefix}/{base}.webp",
                                  "thumb": f"{url_prefix}/{base}-thumb.webp",
                                  "width": w, "height": hh, "sourceFile": f})
            continue
        with Image.open(src_path) as im:
            im = im.convert("RGBA")
            bbox = im.getbbox()  # trim transparent margins
            if bbox:
                im = im.crop(bbox)
            big = im.copy()
            big.thumbnail((max_px, max_px), Image.LANCZOS)
            big.save(out_main, "WEBP", quality=82, method=4)
            small = im.copy()
            small.thumbnail((thumb_px, thumb_px), Image.LANCZOS)
            small.save(out_thumb, "WEBP", quality=80, method=4)
            w, hgt = big.size
        cache[h] = (out_main, out_thumb, w, hgt)
        result[owner].append({
            "angle": angle,
            "src": f"{url_prefix}/{base}.webp",
            "thumb": f"{url_prefix}/{base}-thumb.webp",
            "width": w, "height": hgt,
            "sourceFile": f,
        })
    for m in result:
        result[m].sort(key=lambda x: angle_key(x["angle"]))
    return result, unmatched


# ---------------------------------------------------------------------------
# Last-time-buy (separate Excel)
# ---------------------------------------------------------------------------
def parse_php(v):
    """'PHP 48,310.00' -> 48310 ; numbers pass through ; anything else -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    t = re.sub(r"[^0-9.]", "", str(v))
    if not t:
        return None
    f = float(t)
    return int(f) if f.is_integer() else f


def ltb_cpu_brand(cpu):
    c = (cpu or "").upper()
    if re.search(r"\b(RAI|RYZEN|R[3579]|R[3579]-|AMD)\b", c) or re.match(r"R[3579][ -]", c):
        return "AMD"
    if re.search(r"\b(I[3579]|U[3579]|C[3579]|CORE|ULTRA|INTEL|CELERON|PENTIUM)\b", c) or re.match(r"(I|U|C)[3579]-", c):
        return "Intel"
    if re.search(r"\b(QC|SD|SNAPDRAGON|X1E|X1P|X1)\b", c) or c.startswith("QC-"):
        return "Qualcomm"
    return None


def parse_ltb_specs(spec):
    """Split the short spec string ('RAI 5 330 16GB 512GB 15.6 IPS W11 OPI') into
    processor / memory / storage / the rest. Nothing is rewritten, only split."""
    out = {}
    if not spec:
        return out
    m_ram = re.search(r"\b(\d{1,2}GB)\b((?:\s+(?:D5|D4|DDR5|DDR4|5X|LPDDR5X?|OB))*)", spec)
    if not m_ram:
        out["rest"] = spec
        return out
    cpu = spec[:m_ram.start()].strip()
    out["cpu"] = cpu or None
    out["ram"] = (m_ram.group(1) + m_ram.group(2)).strip()
    after = spec[m_ram.end():]
    m_sto = re.search(r"\b(\d{3,4}GB|\d(?:\.\d)?TB)\b", after)
    if m_sto:
        out["storage"] = m_sto.group(1)
        rest = (after[:m_sto.start()] + " " + after[m_sto.end():]).strip()
    else:
        rest = after.strip()
    rest = re.sub(r"\s+", " ", rest)
    out["rest"] = rest or None
    disp = []
    m_size = re.search(r'\b(1[1-7](?:\.\d)?)(?:"|\b|(?=WUXGA|WQXGA|FHD|OLED))', rest)
    if m_size:
        disp.append(m_size.group(1) + '"')
    for tok in re.findall(r"(?:\b|(?<=\d))(2\.2K|2K|3K|WQXGA|WUXGA|FHD|OLED|IPS|TOUCH|120HZ)\b", rest, re.I):
        disp.append(tok.upper().replace("TOUCH", "Touch").replace("120HZ", "120Hz"))
    if disp:
        out["displayShort"] = " ".join(dict.fromkeys(disp))
    return out


def ltb_availability(qty):
    """Last-time-buy stock is published as a label only - the actual quantity is never written to the site."""
    if not isinstance(qty, (int, float)):
        return None
    if qty <= 0:
        return "Sold out"
    if qty <= 10:
        return "Very limited"
    if qty <= 30:
        return "Limited"
    return "On hand"


def part_from_ltb_model(model):
    """LTB Model text looks like 'HP-NB D40T5PA OMNIBOOK3 15-FN0046AU' - the HP part number is the 2nd word."""
    w = str(model or "").split()
    if len(w) > 1 and w[0].upper().startswith("HP-") and re.fullmatch(r"[A-Z0-9]{6,8}", w[1].upper()):
        return w[1].upper()
    return None


def report_changes(dataset, prev_products, new_products):
    """Log today's differences to data/changes.json (drives the New / Price drop badges)."""
    evs = changes.update(os.path.join(DATA_OUT, "changes.json"), dataset, prev_products, new_products)
    if prev_products is None:
        print(f"[{dataset}] No previous data to compare with - change badges start from the next update.")
        return
    kinds = {}
    for e in evs:
        kinds[e["kind"]] = kinds.get(e["kind"], 0) + 1
    print(f"[{dataset}] Changes recorded today: " + (", ".join(f"{v} {k}" for k, v in kinds.items()) or "none"))


def build_ltb(excel_path, images_dir=None):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip().lower() == "model" for c in r if c))
    headers = [clean(h) for h in rows[hi]]
    recs = []
    for r in rows[hi + 1:]:
        rec = {str(h): clean(v) for h, v in zip(headers, r) if h}
        if rec.get("Model"):
            recs.append(rec)
    seen, uniq, dups = set(), [], []
    for r in recs:
        if r["Model"] in seen:
            dups.append(r["Model"]); continue
        seen.add(r["Model"]); uniq.append(r)
    if dups:
        print("WARNING duplicate LTB models skipped:", dups)

    images = {}
    if images_dir:
        images, unmatched = build_images([r["Model"] for r in uniq], images_dir,
                                         out_dir=os.path.join(ROOT, "assets", "images", "ltb"),
                                         url_prefix="assets/images/ltb")
        if unmatched:
            print("LTB image files NOT matched (skipped):", unmatched)
    else:
        prev = os.path.join(DATA_OUT, "ltb.json")
        if os.path.exists(prev):
            with open(prev, encoding="utf-8") as fh:
                for p in json.load(fh).get("products", []):
                    images[p["model"]] = p.get("images", [])

    out = []
    for i, r in enumerate(uniq):
        tag = r.get("Tagging") or ""
        seg, _, cat = tag.partition(" - ")
        spec = parse_ltb_specs(r.get("Specs"))
        sku = sku_from_model(r["Model"])
        p = {
            "id": "ltb-" + slug(sku),
            "model": r["Model"],
            "sku": sku,
            "partNo": part_from_ltb_model(r["Model"]),
            "matNo": str(r["Mat No."]) if r.get("Mat No.") is not None else None,
            "tagging": tag or None,
            "segment": seg.strip() if cat else None,
            "type": (cat or tag).strip() or "Other",
            "cpu": spec.get("cpu"),
            "cpuBrand": ltb_cpu_brand(spec.get("cpu")),
            "ram": spec.get("ram"),
            "storage": spec.get("storage"),
            "displayShort": spec.get("displayShort"),
            "otherSpecs": [spec["rest"]] if spec.get("rest") else None,
            "specsRaw": r.get("Specs"),
            "dp": parse_php(r.get("DP")),
            "srp": parse_php(r.get("SRP")),
            "sale": parse_php(r.get("Special Price")),
            "availability": ltb_availability(r.get("Qty")) or r.get("Availability"),
            "priceText": {"dp": r.get("DP"), "srp": r.get("SRP"), "sale": r.get("Special Price")},
            "images": images.get(r["Model"], []),
            "row": i + 1,
        }
        out.append({k: v for k, v in p.items() if v not in (None, [], "")})

    with open(excel_path, "rb") as fh:
        digest = hashlib.sha1(fh.read()).hexdigest()[:10]
    payload = {"generatedAt": datetime.now().isoformat(timespec="seconds"),
               "sourceFile": os.path.basename(excel_path), "sourceHash": digest,
               "currency": "PHP", "count": len(out), "products": out}
    ltb_path = os.path.join(DATA_OUT, "ltb.json")
    prev_products = changes.load_previous(ltb_path)
    with open(ltb_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    report_changes("ltb", prev_products, out)
    no_img = [p["model"] for p in out if not p.get("images")]
    print(f"Wrote {len(out)} last-time-buy products." + (f" Without photos (silhouette shown): {no_img}" if no_img else ""))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", default=os.path.join(ROOT, "source", "Current Pricelist - HPNB.xlsx"))
    ap.add_argument("--images", default=None, help="Folder of product photos named '<Model> - <Angle>.png'")
    ap.add_argument("--ltb-excel", default=os.path.join(ROOT, "source", "Last-time-buy models.xlsx"))
    ap.add_argument("--ltb-images", default=None, help="Folder of Last-time-buy photos named '<Model> - <Angle>.png'")
    ap.add_argument("--only-ltb", action="store_true", help="Only rebuild data/ltb.json")
    ap.add_argument("--specsheets", default=os.path.join(ROOT, "source", "specsheets"),
                    help="Folder of HP technical spec sheets (<PartNo>_*.xlsx) to add full specs + security features")
    args = ap.parse_args()
    if os.path.exists(args.ltb_excel):
        build_ltb(args.ltb_excel, args.ltb_images)
    else:
        print("No Last-time-buy Excel found at", args.ltb_excel, "(skipped)")
    if args.only_ltb:
        return

    headers, rows = read_excel(args.excel)
    print(f"Read {len(rows)} rows from {os.path.basename(args.excel)}; columns: {[h for h in headers if h]}")

    # de-duplicate by exact model (keep first occurrence, report the rest)
    seen, products, dups = set(), [], []
    for r in rows:
        if r["Model"] in seen:
            dups.append(r["Model"])
            continue
        seen.add(r["Model"])
        products.append(r)
    if dups:
        print("WARNING duplicate models skipped:", dups)

    for r in products:
        pn = r.get("Part Number")
        if pn is not None and not PART_RE.match(str(pn).strip().upper()):
            print(f"WARNING Part Number '{pn}' for '{r['Model']}' does not look like an HP part number - please check the Excel.")
    part_alias = {str(r["Part Number"]).strip().upper(): r["Model"] for r in products
                  if r.get("Part Number") is not None and PART_RE.match(str(r["Part Number"]).strip().upper())}
    sheets = load_specsheets(args.specsheets)
    if sheets:
        print("Spec sheets:", ", ".join(f"{k} ({'matched' if k in part_alias else 'NO pricelist row'})" for k in sheets))

    images = {}
    unmatched = []
    # previously generated image list (kept for any model with no photo in --images)
    prev_images = {}
    prev = os.path.join(DATA_OUT, "products.json")
    if os.path.exists(prev):
        with open(prev, encoding="utf-8") as fh:
            for p in json.load(fh).get("products", []):
                prev_images[p["model"]] = p.get("images", [])
    if args.images:
        images, unmatched = build_images([p["Model"] for p in products], args.images, aliases=part_alias)
        for m, imgs in prev_images.items():
            if not images.get(m) and imgs:
                images[m] = imgs
    else:
        images = prev_images

    # type info (keep any edits made by hand in data/types.json)
    types_path = os.path.join(DATA_OUT, "types.json")
    type_info = {}
    if os.path.exists(types_path):
        with open(types_path, encoding="utf-8") as fh:
            type_info = {t["code"]: t for t in json.load(fh)}

    known_cols = {"Platform", "Type", "Model", "Specs", "SRP", "DP", "Promo SRP", "Promo DP", "Qty", "Status", "Comments", "Part Number", "Material Number"}
    out, type_order = [], []
    for i, r in enumerate(products):
        spec = parse_specs(r.get("Specs"))
        norm = lambda x: re.sub(r"[^a-z0-9]", "", str(x).lower())
        if spec.get("other") and r.get("Platform"):
            # the platform codename repeated inside Specs is already shown from the Platform column
            spec["other"] = [o for o in spec["other"] if not (norm(o).startswith(norm(r["Platform"])) or norm(r["Platform"]).startswith(norm(o)))] or None
        t = r.get("Type") or "Other"
        if t not in type_order:
            type_order.append(t)
        srp, dp = r.get("SRP"), r.get("DP")
        pn_ok = str(r["Part Number"]).strip().upper() if r.get("Part Number") is not None and PART_RE.match(str(r["Part Number"]).strip().upper()) else None
        sku = product_sku(r["Model"], pn_ok)
        sheet = sheets.get(pn_ok) if pn_ok else None
        if sheet:
            # pricelist Specs text that is not in the usual pipe format is kept as raw text only;
            # the card fields come from HP's own spec sheet
            if "|" not in (r.get("Specs") or ""):
                spec = {}
            w = re.search(r"\b(\d/\d/\d)\b", r.get("Specs") or "")
            if w and not spec.get("warranty"):
                spec["warranty"] = w.group(1)
            for k, v in sheet["_fill"].items():
                if v and not spec.get(k):
                    spec[k] = v
        p = {
            "id": slug(sku),
            "model": r["Model"],
            "title": sheet["name"] if sheet else None,
            "sku": sku,
            "partNo": str(r["Part Number"]).strip() if r.get("Part Number") is not None else None,
            "matNo": str(r["Material Number"]).strip() if r.get("Material Number") is not None else None,
            "type": t,
            "platform": r.get("Platform"),
            "srp": srp if isinstance(srp, (int, float)) else None,
            "dp": dp if isinstance(dp, (int, float)) else None,
            "promoSrp": r.get("Promo SRP") if isinstance(r.get("Promo SRP"), (int, float)) else None,
            "promoDp": r.get("Promo DP") if isinstance(r.get("Promo DP"), (int, float)) else None,
            # "Comments" column, e.g. "Maximum 50 units only until Oct 31, 2026" - shown with the promo price
            "note": str(r["Comments"]).strip() if r.get("Comments") not in (None, "") else None,
            "status": str(r["Status"]).strip() if r.get("Status") not in (None, "") else None,
            "cpu": spec.get("cpu"),
            "cpuBrand": cpu_brand(spec.get("cpu")),
            "ram": spec.get("ram"),
            "storage": spec.get("storage"),
            "graphics": spec.get("graphics"),
            "display": spec.get("display"),
            "displayShort": spec.get("displayShort") or display_short(spec.get("display")),
            "os": spec.get("os"),
            "color": spec.get("color"),
            "camera": spec.get("camera"),
            "formFactor": spec.get("formFactor"),
            "series": spec.get("series"),
            "adapter": spec.get("adapter"),
            "warranty": spec.get("warranty"),
            "carePack": spec.get("carePack"),
            "otherSpecs": spec.get("other"),
            "specsRaw": r.get("Specs"),
            "images": images.get(r["Model"], []),
            "specSheet": {k: v for k, v in sheet.items() if not k.startswith("_")} if sheet else None,
            "row": i + 1,
        }
        # Comments that change how a promo price is used:
        #  "Don't offer yet - for reference only"  -> shown as an upcoming promo, regular price stays the selling price
        #  "No available promo priced units ..."    -> promo units used up, regular price shown
        note_l = (p.get("note") or "").lower().replace("’", "'")
        if (p.get("promoSrp") is not None or p.get("promoDp") is not None) and re.search(r"don'?t offer yet|reference only", note_l):
            p["promoState"] = "upcoming"
            p["upcomingSrp"], p["upcomingDp"] = p.pop("promoSrp"), p.pop("promoDp")
        elif (p.get("promoSrp") is not None or p.get("promoDp") is not None) and re.search(r"no available promo|promo[- ]priced units? (are )?(sold out|used up|not available)", note_l):
            p["promoState"] = "soldout"
            p["endedSrp"], p["endedDp"] = p.pop("promoSrp"), p.pop("promoDp")
        # extra, unknown columns are carried through (never invented)
        extra = {k: v for k, v in r.items() if k not in known_cols and v is not None}
        if extra:
            p["extra"] = extra
        out.append({k: v for k, v in p.items() if v not in (None, [], "")})

    rank = {c.lower(): i for i, c in enumerate(TYPE_RANK)}
    first_seen = list(type_order)
    type_order.sort(key=lambda c: (rank.get(c.lower(), len(rank)), first_seen.index(c)))
    types = []
    for code in type_order:
        prev = type_info.get(code)
        if prev:
            types.append(prev)
        else:
            label, desc = TYPE_INFO_DEFAULTS.get(code, (code, f"HP {code} models currently on the pricelist."))
            types.append({"code": code, "label": label, "description": desc})

    os.makedirs(DATA_OUT, exist_ok=True)
    with open(args.excel, "rb") as fh:
        digest = hashlib.sha1(fh.read()).hexdigest()[:10]
    payload = {
        "generatedAt": datetime.now().isoformat(timespec="seconds"),
        "sourceFile": os.path.basename(args.excel),
        "sourceHash": digest,
        "currency": "PHP",
        "count": len(out),
        "products": out,
    }
    products_path = os.path.join(DATA_OUT, "products.json")
    prev_products = changes.load_previous(products_path)
    with open(products_path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, ensure_ascii=False, indent=2)
    report_changes("main", prev_products, out)
    with open(types_path, "w", encoding="utf-8") as fh:
        json.dump(types, fh, ensure_ascii=False, indent=2)

    no_img = [p["model"] for p in out if not p.get("images")]
    print(f"Wrote {len(out)} products, {len(types)} types.")
    if no_img:
        print("Products without images:", no_img)
    if unmatched:
        print("Image files NOT matched to any model (skipped):", unmatched)


if __name__ == "__main__":
    main()
