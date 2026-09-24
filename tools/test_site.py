"""End-to-end checks for the pricelist site (run: python tools/test_site.py http://localhost:8765/).
Requires: pip install playwright openpyxl && playwright install chromium"""
import json, os, re, sys
import openpyxl
from playwright.sync_api import sync_playwright

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://localhost:8765/"
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SHOTS = os.path.join(ROOT, "tools", "screenshots")
os.makedirs(SHOTS, exist_ok=True)
results = []

def check(name, ok, info=""):
    results.append((name, bool(ok), info))
    print(("PASS " if ok else "FAIL ") + name + (f"  [{info}]" if info else ""))

# ---------- data vs Excel ----------
wb = openpyxl.load_workbook(os.path.join(ROOT, "source", "Current Pricelist - HPNB.xlsx"), data_only=True)
ws = wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))
hi = next(i for i, r in enumerate(rows) if r and "Model" in r)
hdr = rows[hi]
xl = [dict(zip(hdr, r)) for r in rows[hi + 1:] if r and r[hdr.index("Model")]]
data = json.load(open(os.path.join(ROOT, "data", "products.json"), encoding="utf-8"))["products"]
by_model = {p["model"]: p for p in data}
check("every Excel model present", all(r["Model"].strip() in by_model for r in xl), f"{len(xl)} rows / {len(data)} products")
check("no duplicate models", len(by_model) == len(data))
check("SRP matches Excel", all(by_model[r["Model"].strip()]["srp"] == r["SRP"] for r in xl))
check("DP matches Excel", all(by_model[r["Model"].strip()]["dp"] == r["DP"] for r in xl))
check("Type matches Excel", all(by_model[r["Model"].strip()]["type"] == r["Type"].strip() for r in xl))
img_ok = all(all(im["sourceFile"].startswith(p["model"] + " - ") for im in p["images"]) for p in data)
check("images only assigned by exact model name", img_ok)
check("every image file exists", all(os.path.exists(os.path.join(ROOT, im[k])) for p in data for im in p["images"] for k in ("src", "thumb")))
check("all asset paths relative", all(not re.match(r"^(/|https?:)", im["src"]) for p in data for im in p["images"]))

with sync_playwright() as pw:
    b = pw.chromium.launch()
    ctx = b.new_context(viewport={"width": 1440, "height": 900})
    pg = ctx.new_page()
    bad = []
    pg.on("response", lambda r: bad.append(r.url) if r.status >= 400 else None)
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto(BASE); pg.wait_for_selector(".card")
    count = lambda: pg.locator("#products > article").count()
    check("renders all products", count() == len(data), str(count()))
    pg.screenshot(path=f"{SHOTS}/desktop-grid.png", full_page=False)

    def search(q):
        pg.fill("#search", q); pg.wait_for_timeout(250)
        return [pg.locator("#products > article").nth(i).get_attribute("data-id") for i in range(count())]
    r = search("14-kb0105TU"); check("search exact SKU", r == ["14-kb0105tu"], str(r))
    r = search("kb0105"); check("search partial SKU", r == ["14-kb0105tu"], str(r))
    r = search("Ultra 7"); exp = sorted(p["id"] for p in data if "Ultra 7" in (p.get("cpu") or ""))
    check("search 'Ultra 7'", sorted(r) == exp and len(exp) > 0, f"{len(r)} results")
    r = search("ryzen"); check("search 'ryzen'", len(r) == sum(1 for p in data if p.get("cpuBrand") == "AMD"), str(len(r)))
    r = search("OmniBook 5 Flip"); check("search product name", len(r) == 1, str(r))
    r = search("zzzz-nothing"); check("empty state shown", pg.is_visible("#emptyState") and r == [])
    search("")

    # filters
    def tick(group, val, on=True):
        box = pg.locator(f'#filters input[data-group="{group}"][value="{val}"]')
        if box.is_checked() != on:
            box.locator("xpath=..").click()
        pg.wait_for_timeout(120)
    tick("cpu", "Intel")
    exp = [p for p in data if p.get("cpuBrand") == "Intel"]
    check("CPU Intel filter", count() == len(exp), f"{count()} vs {len(exp)}")
    tick("cpu", "Intel", False); tick("cpu", "AMD")
    exp = [p for p in data if p.get("cpuBrand") == "AMD"]
    check("CPU AMD filter", count() == len(exp), f"{count()} vs {len(exp)}")
    tick("cpu", "AMD", False)
    tick("price", "60000")
    exp = [p for p in data if 60000 <= p["srp"] <= 69999]
    check("price band 60k-69,999", count() == len(exp), f"{count()} vs {len(exp)}")
    tick("price", "60000", False); tick("price", "100000plus")
    exp = [p for p in data if p["srp"] >= 100000]
    check("price band 100k+", count() == len(exp), f"{count()} vs {len(exp)}")
    tick("price", "100000plus", False)
    tick("type", "OB3"); tick("cpu", "Intel"); tick("price", "60000")
    exp = [p for p in data if p["type"] == "OB3" and p.get("cpuBrand") == "Intel" and 60000 <= p["srp"] <= 69999]
    check("combined Type+CPU+Price", count() == len(exp) and len(exp) > 0, f"{count()} vs {len(exp)}")
    check("result count text", str(len(exp)) in pg.inner_text("#resultCount"))
    pg.click("#clearFilters"); pg.wait_for_timeout(150)
    check("clear filters", count() == len(data))

    # type tabs
    types = json.load(open(os.path.join(ROOT, "data", "types.json"), encoding="utf-8"))
    for t in types:
        pg.click(f'.nav-link[data-nav="{t["code"]}"]'); pg.wait_for_timeout(150)
        n = sum(1 for p in data if p["type"] == t["code"])
        ok = count() == n and pg.inner_text("#catTitle") == t["label"]
        check(f"type tab {t['code']}", ok, f"{count()} vs {n}")
    pg.click('.nav-link[data-nav="all"]'); pg.wait_for_timeout(150)

    # view switch keeps filters, remembered in session
    pg.fill("#search", "OmniBook X"); pg.wait_for_timeout(200)
    n_before = count()
    pg.click('.vt-btn[data-view="list"]'); pg.wait_for_timeout(150)
    check("list view keeps search", pg.locator("#products.list .row").count() == n_before)
    pg.screenshot(path=f"{SHOTS}/desktop-list.png")
    pg.reload(); pg.wait_for_selector("#products > article")
    check("view remembered in session", "list" in pg.get_attribute("#products", "class"))
    pg.click('.vt-btn[data-view="grid"]'); pg.fill("#search", ""); pg.wait_for_timeout(150)

    # details + gallery
    pg.locator('.card [data-open="14-kb0100tu"].btn').click()
    pg.wait_for_selector("#detailModal:not([hidden])")
    check("detail modal title", pg.inner_text("#dmTitle") == by_model["HP OmniBook X Flip NG AI PC 14-kb0100TU"]["model"])
    txt = pg.inner_text("#detailBody")
    check("detail shows SRP/DP", "₱112,990" in txt and "₱90,392" in txt)
    src0 = pg.get_attribute("#gMain img", "src")
    pg.locator(".thumb").nth(2).click(); pg.wait_for_timeout(100)
    src2 = pg.get_attribute("#gMain img", "src")
    check("thumbnail switches main image", src0 != src2 and "rear-left" in src2, src2)
    pg.screenshot(path=f"{SHOTS}/desktop-detail.png")
    pg.click("#gZoom"); pg.wait_for_selector("#lightbox:not([hidden])")
    check("lightbox opens", pg.is_visible("#lbImg"))
    pg.keyboard.press("ArrowRight"); pg.wait_for_timeout(80)
    check("lightbox next", "left-profile" in pg.get_attribute("#lbImg", "src"))
    pg.keyboard.press("Escape"); pg.keyboard.press("Escape"); pg.wait_for_timeout(100)
    check("modal closes", pg.is_hidden("#detailModal"))
    pg.goto(BASE + "#/type/OBX/product/14-kc0081au"); pg.wait_for_selector("#detailModal:not([hidden])")
    check("deep link to product", "14-kc0081AU" in pg.inner_text("#dmTitle"))
    pg.keyboard.press("Escape")

    # sites
    pg.click('.nav-link[data-nav="sites"]'); pg.wait_for_selector("#sitesView:not([hidden])")
    links = pg.eval_on_selector_all(".site-card a", "els => els.map(e => [e.href, e.target, e.rel])")
    exp_urls = ["https://pcb.inc.hp.com/webapp/#/ap-en?hierarchy=F&status=L&status=O",
                "https://support.hp.com/ph-en/help/service-center",
                "https://cpc2.ext.hp.com/?countries=PH"]
    check("related site URLs", [l[0] for l in links] == exp_urls, str(links))
    check("links open new tab", all(l[1] == "_blank" and "noopener" in l[2] for l in links))
    pg.screenshot(path=f"{SHOTS}/desktop-sites.png")


    # ---------- Promos ----------
    promos = json.load(open(os.path.join(ROOT, "data", "promos.json"), encoding="utf-8"))
    pg.goto(BASE + "#/promos"); pg.wait_for_selector(".promo-card")
    check("promos tab shows both promos", pg.locator(".promo-card").count() == 2)
    txt = pg.inner_text("#promosGrid")
    check("promo period shown", txt.count("August 1 – October 31, 2026") == 2, "")
    gl = pg.get_attribute("#promo-gcash .promo-actions a", "href")
    check("GCash links to official page", gl == "https://h41201.www4.hp.com/WMCF.Web/ph/en/promotion/16376/Home/" and pg.get_attribute("#promo-gcash .promo-actions a", "target") == "_blank")
    check("Home Credit has no external link", pg.locator("#promo-home-credit a[href^='http']").count() == 0)
    pg.click("#promo-home-credit .promo-actions button"); pg.wait_for_selector("#lightbox:not([hidden])")
    check("Home Credit flyer opens enlarged", "home-credit" in pg.get_attribute("#lbImg", "src"))
    pg.keyboard.press("Escape")
    check("freebie banner shown", "A08JTAA" in txt)
    pg.evaluate("window._opened = []; window.open = (u, t) => { window._opened.push([u, t]); }")
    pg.locator("#promo-gcash .promo-caption").click()
    op = pg.evaluate("window._opened")
    check("whole GCash card opens promo in new tab", any(o and o[0] and "promotion/16376" in o[0] and o[1] == "_blank" for o in op), str(op))
    pg.evaluate("window._opened = []"); pg.locator("#promo-home-credit .promo-caption").click()
    check("Home Credit card is not a link", not any(o and o[0] for o in pg.evaluate("window._opened")))
    elig = {e["sku"]: e["reward"] for e in promos["promos"][0]["eligible"] if e["sku"]}
    pg.goto(BASE + "#/"); pg.wait_for_selector(".card")
    bad_badge = []
    for p in data:
        want = elig.get(p["sku"].upper())
        loc = pg.locator(f'.card[data-id="{p["id"]}"] .promo-badge')
        got = loc.inner_text().strip() if loc.count() else None
        if (want is None) != (got is None) or (want and f"₱{want:,}" not in got):
            bad_badge.append((p["sku"], want, got))
    check("GCash badges match official eligible list", not bad_badge, str(bad_badge[:4]))
    check("freebie shown on every laptop card", pg.locator(".card .freebie-line").count() == len(data))

    # ---------- Last-time-buy ----------
    wb2 = openpyxl.load_workbook(os.path.join(ROOT, "source", "Last-time-buy models.xlsx"), data_only=True)
    r2 = list(wb2.worksheets[0].iter_rows(values_only=True)); h2 = r2[0]
    xl2 = [dict(zip(h2, r)) for r in r2[1:] if r and r[h2.index("Model")]]
    ltb = json.load(open(os.path.join(ROOT, "data", "ltb.json"), encoding="utf-8"))["products"]
    num = lambda v: float(re.sub(r"[^0-9.]", "", str(v)))
    by2 = {p["model"]: p for p in ltb}
    check("LTB: every Excel model present, no duplicates", len(by2) == len(ltb) == len(xl2) and all(r["Model"] in by2 for r in xl2))
    check("LTB: sale/DP/SRP/Qty match Excel", all(num(r["Special Price"]) == by2[r["Model"]]["sale"] and num(r["DP"]) == by2[r["Model"]]["dp"]
          and num(r["SRP"]) == by2[r["Model"]]["srp"] and r["Qty"] == by2[r["Model"]]["qty"] for r in xl2))
    check("LTB: specs text preserved", all(by2[r["Model"]]["specsRaw"] == r["Specs"] for r in xl2))
    check("LTB: images only by exact model name", all(im["sourceFile"].startswith(p["model"] + " - ") for p in ltb for im in p.get("images", [])))
    pg.goto(BASE + "#/ltb"); pg.wait_for_selector(".ltb-card")
    check("LTB tab renders all models", pg.locator(".ltb-card").count() == len(ltb))
    check("LTB: badge + sale price on every card", pg.locator(".ltb-card .ltb-badge").count() == len(ltb) and pg.locator(".ltb-card .sale-box").count() == len(ltb))
    first = pg.locator(".ltb-card").first
    check("LTB: first sale price visible without scrolling (desktop)", first.locator(".sale-box").bounding_box()["y"] < 900)
    fs = pg.eval_on_selector(".ltb-card .sale-box .val", "e => parseFloat(getComputedStyle(e).fontSize)")
    fdp = pg.eval_on_selector(".ltb-card .was", "e => parseFloat(getComputedStyle(e).fontSize)")
    check("LTB: sale price is the most prominent price", fs > fdp * 1.6, f"{fs} vs {fdp}")
    tick("cpu", "AMD")
    check("LTB: CPU filter", count() == sum(1 for p in ltb if p.get("cpuBrand") == "AMD"))
    tick("cpu", "AMD", False); tick("type", "Pav")
    check("LTB: type filter", count() == sum(1 for p in ltb if p["type"] == "Pav"))
    tick("type", "Pav", False); tick("price", "50000")
    check("LTB: price band uses sale price", count() == sum(1 for p in ltb if 50000 <= p["sale"] < 60000))
    pg.click("#clearFilters"); pg.wait_for_timeout(100)
    r = search("fp0061"); check("LTB: search", r == ["ltb-14-fp0061tu"], str(r)); search("")
    pg.click('.vt-btn[data-view="list"]'); pg.wait_for_timeout(100)
    check("LTB: list view", pg.locator("#products.list .ltb-row").count() == len(ltb))
    pg.click('.vt-btn[data-view="grid"]')
    pg.locator('.ltb-card[data-id="ltb-14-fe0028qu"] .btn').click(); pg.wait_for_selector("#detailModal:not([hidden])")
    dt = pg.inner_text("#detailBody")
    check("LTB detail: sale, DP, SRP, stock, material no.", all(x in dt for x in ["₱84,021", "₱86,620", "₱108,990", "19 units", "432115035942"]))
    check("LTB detail: silhouette when no photo", "laptop-placeholder" in pg.get_attribute("#gallery img", "src"))
    pg.keyboard.press("Escape")
    pg.goto(BASE + "#/ltb/product/ltb-13-bg1055au"); pg.wait_for_selector("#detailModal:not([hidden])")
    check("LTB: deep link + gallery", pg.locator(".thumb").count() == 4)
    pg.keyboard.press("Escape")
    pg.click('.nav-link[data-nav="all"]'); pg.wait_for_timeout(200)
    check("back to main list after LTB", count() == len(data) and pg.locator(".ltb-card").count() == 0)
    bad2 = pg.eval_on_selector_all("img", "els => els.filter(e => e.complete && e.naturalWidth === 0).map(e => e.src)")
    check("no broken images (after new tabs)", not bad2, str(bad2[:3]))

    # broken images on all product images
    pg.goto(BASE + "#/"); pg.wait_for_selector(".card")
    pg.evaluate("window.scrollTo(0, document.body.scrollHeight)"); pg.wait_for_timeout(800)
    broken = pg.eval_on_selector_all("img", "els => els.filter(e => e.complete && e.naturalWidth === 0).map(e => e.src)")
    check("no broken images", not broken, str(broken[:3]))
    check("no HTTP errors", not bad, str(bad[:3]))
    check("no JS errors", not errs, str(errs[:3]))
    abs_refs = pg.evaluate("""() => [...document.querySelectorAll('link[href],script[src],img[src]')]
        .map(e => e.getAttribute('href') || e.getAttribute('src')).filter(u => /^\\//.test(u))""")
    check("no root-absolute asset paths", not abs_refs, str(abs_refs))

    # mobile
    m = b.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True, device_scale_factor=2)
    mp = m.new_page(); merr = []
    mp.on("pageerror", lambda e: merr.append(str(e)))
    mp.goto(BASE); mp.wait_for_selector(".card")
    ow = mp.evaluate("document.documentElement.scrollWidth > window.innerWidth")
    check("mobile: no horizontal scroll", not ow)
    check("mobile: nav collapsed", not mp.is_visible("#primaryNav"))
    check("mobile: search full width", mp.evaluate("document.querySelector('.search').getBoundingClientRect().width") > 340)
    mp.screenshot(path=f"{SHOTS}/mobile-grid.png")
    mp.click("#menuToggle"); mp.wait_for_timeout(150)
    check("mobile: hamburger opens nav", mp.is_visible('.nav-link[data-nav="OB5"]'))
    mp.screenshot(path=f"{SHOTS}/mobile-menu.png")
    mp.click('.nav-link[data-nav="OB5"]'); mp.wait_for_timeout(200)
    check("mobile: nav closes after pick", not mp.is_visible("#primaryNav") and mp.locator("#products > article").count() == 2)
    mp.click("#filtersToggle"); mp.wait_for_timeout(350)
    check("mobile: filter drawer opens", mp.is_visible("#filtersApply"))
    mp.screenshot(path=f"{SHOTS}/mobile-filters.png")
    mp.click("#filtersApply"); mp.wait_for_timeout(300)
    mp.click('.nav-link[data-nav="all"]', force=True) if mp.is_visible('.nav-link[data-nav="all"]') else mp.goto(BASE + "#/")
    mp.wait_for_timeout(200)
    mp.click('.vt-btn[data-view="list"]'); mp.wait_for_timeout(150)
    mp.screenshot(path=f"{SHOTS}/mobile-list.png")
    check("mobile list: no horizontal scroll", not mp.evaluate("document.documentElement.scrollWidth > window.innerWidth"))
    mp.locator("#products .row .btn").first.click(); mp.wait_for_selector("#detailModal:not([hidden])")
    pw_ = mp.evaluate("document.querySelector('.modal-panel').scrollWidth <= document.querySelector('.modal-panel').clientWidth")
    check("mobile: detail fits width", pw_)
    s0 = mp.get_attribute("#gMain img", "src")
    box = mp.locator("#gMain").bounding_box()
    y = box["y"] + box["height"] / 2
    mp.evaluate(f"""() => {{
        const el = document.querySelector('#gMain');
        const mk = (x) => new Touch({{identifier: 1, target: el, clientX: x, clientY: {y}}});
        el.dispatchEvent(new TouchEvent('touchstart', {{touches: [mk(300)], changedTouches: [mk(300)], bubbles: true}}));
        el.dispatchEvent(new TouchEvent('touchend', {{touches: [], changedTouches: [mk(120)], bubbles: true}}));
    }}""")
    mp.wait_for_timeout(100)
    check("mobile: swipe changes image", mp.get_attribute("#gMain img", "src") != s0)
    mp.screenshot(path=f"{SHOTS}/mobile-detail.png")
    mp.locator(".modal-panel").evaluate("e => e.scrollTop = 600"); mp.wait_for_timeout(100)
    mp.screenshot(path=f"{SHOTS}/mobile-detail-specs.png")

    mp.goto(BASE + "#/ltb"); mp.wait_for_selector(".ltb-row, .ltb-card"); mp.keyboard.press("Escape")
    mp.click('.vt-btn[data-view="grid"]'); mp.wait_for_selector(".ltb-card"); mp.evaluate("window.scrollTo(0,0)"); mp.wait_for_timeout(300)
    sb = mp.locator(".ltb-card .sale-box").first.bounding_box()
    check("mobile LTB: sale price visible on first screen", sb["y"] + sb["height"] < 844, str(sb))
    check("mobile LTB: no horizontal scroll", not mp.evaluate("document.documentElement.scrollWidth > window.innerWidth"))
    mp.screenshot(path=f"{SHOTS}/mobile-ltb.png")
    mp.goto(BASE + "#/promos"); mp.wait_for_selector(".promo-card"); mp.wait_for_timeout(300)
    check("mobile promos: no horizontal scroll", not mp.evaluate("document.documentElement.scrollWidth > window.innerWidth"))
    cta = mp.locator("#promo-gcash .promo-actions .btn").bounding_box()
    check("mobile promos: CTA easy to tap", cta["height"] >= 44 and cta["width"] > 300, str(cta))
    mp.screenshot(path=f"{SHOTS}/mobile-promos.png", full_page=True)
    mp.click("#menuToggle"); mp.wait_for_timeout(150)
    check("mobile: menu has Promos + Last-time-buy", mp.is_visible('.nav-link[data-nav="promos"]') and mp.is_visible('.nav-link[data-nav="ltb"]'))
    mp.click('.nav-link[data-nav="ltb"]'); mp.wait_for_timeout(200)
    check("mobile: nav to LTB works", mp.locator(".ltb-card").count() > 0)
    small = mp.eval_on_selector_all(".btn, .vt-btn, .menu-toggle, .thumb", "els => els.filter(e => e.offsetParent && e.getBoundingClientRect().height < 36).map(e => e.className)")
    check("mobile: touch targets >= 36px", not small, str(small[:4]))
    check("mobile: no JS errors", not merr, str(merr[:3]))
    t = b.new_context(viewport={"width": 820, "height": 1180}).new_page()
    t.goto(BASE); t.wait_for_selector(".card"); t.screenshot(path=f"{SHOTS}/tablet-grid.png")
    check("tablet: no horizontal scroll", not t.evaluate("document.documentElement.scrollWidth > window.innerWidth"))
    b.close()

failed = [r for r in results if not r[1]]
print(f"\n{len(results) - len(failed)}/{len(results)} checks passed")
sys.exit(1 if failed else 0)
